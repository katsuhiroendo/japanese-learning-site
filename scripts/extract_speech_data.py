#!/usr/bin/env python3
"""
スライド資料から日本語発音データを抽出し、js/speech-data.js を生成するスクリプト
"""

import json
import os
import re
from pptx import Presentation

# 基本モーラ（ローマ字 -> ひらがな）マップ
MORA_MAP = {
    'a': 'あ', 'i': 'い', 'u': 'う', 'e': 'え', 'o': 'お',
    'ka': 'か', 'ki': 'き', 'ku': 'く', 'ke': 'け', 'ko': 'こ',
    'sa': 'さ', 'shi': 'し', 'su': 'す', 'se': 'せ', 'so': 'そ',
    'ta': 'た', 'chi': 'ち', 'tsu': 'つ', 'te': 'て', 'to': 'と',
    'na': 'な', 'ni': 'に', 'nu': 'ぬ', 'ne': 'ね', 'no': 'の',
    'ha': 'は', 'hi': 'ひ', 'fu': 'ふ', 'he': 'へ', 'ho': 'ほ',
    'ma': 'ま', 'mi': 'み', 'mu': 'む', 'me': 'め', 'mo': 'も',
    'ya': 'や', 'yu': 'ゆ', 'yo': 'よ',
    'ra': 'ら', 'ri': 'り', 'ru': 'る', 're': 'れ', 'ro': 'ろ',
    'wa': 'わ', 'wo': 'を', 'n': 'ん',
    'ga': 'が', 'gi': 'ぎ', 'gu': 'ぐ', 'ge': 'げ', 'go': 'ご',
    'za': 'ざ', 'ji': 'じ', 'zu': 'ず', 'ze': 'ぜ', 'zo': 'ぞ',
    'da': 'だ', 'de': 'で', 'do': 'ど',
    'ba': 'ば', 'bi': 'び', 'bu': 'ぶ', 'be': 'べ', 'bo': 'ぼ',
    'pa': 'ぱ', 'pi': 'ぴ', 'pu': 'ぷ', 'pe': 'ぺ', 'po': 'ぽ',
    'kya': 'きゃ', 'kyu': 'きゅ', 'kyo': 'きょ',
    'sha': 'しゃ', 'shu': 'しゅ', 'sho': 'しょ',
    'cha': 'ちゃ', 'chu': 'ちゅ', 'cho': 'ちょ',
    'nya': 'にゃ', 'nyu': 'にゅ', 'nyo': 'にょ',
    'hya': 'ひゃ', 'hyu': 'ひゅ', 'hyo': 'ひょ',
    'mya': 'みゃ', 'myu': 'みゅ', 'myo': 'みょ',
    'rya': 'りゃ', 'ryu': 'りゅ', 'ryo': 'りょ',
    'gya': 'ぎゃ', 'gyu': 'ぎゅ', 'gyo': 'ぎょ',
    'ja': 'じゃ', 'ju': 'じゅ', 'jo': 'じょ',
    'bya': 'びゃ', 'byu': 'びゅ', 'byo': 'びょ',
    'pya': 'ぴゃ', 'pyu': 'ぴゅ', 'pyo': 'ぴょ',
}

def mora_to_kana(m):
    m = m.lower().strip()
    if m in MORA_MAP:
        return MORA_MAP[m]
    # 促音 (っ) の処理: tto, kki, ssu, etc.
    if len(m) >= 2 and m[0] == m[1] and m[0] not in 'aeioun':
        rest = m[1:]
        if rest in MORA_MAP:
            return 'っ' + MORA_MAP[rest]
    return m

def extract_001_fc(filepath):
    prs = Presentation(filepath)
    data = {}
    for i, slide in enumerate(prs.slides):
        texts = [shape.text.strip() for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]
        if not texts:
            continue
        romaji = texts[0]
        meaning = texts[1] if len(texts) > 1 else ''
        moras = texts[2:]
        kana = ''.join(mora_to_kana(m) for m in moras) if moras else romaji
        data[i + 1] = {
            'romaji': romaji,
            'kana': kana,
            'en': meaning
        }
    return data

def extract_002_fc(filepath):
    prs = Presentation(filepath)
    data = {}
    for i, slide in enumerate(prs.slides):
        texts = [shape.text.strip() for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]
        if not texts:
            continue

        # ヘッダー行を除去
        content_texts = [t for t in texts if not t.startswith('Lesson 1')]
        if not content_texts:
            data[i + 1] = {
                'romaji': 'AISATSU',
                'kana': 'あいさつ',
                'en': 'Greetings'
            }
            continue

        romaji = content_texts[0]
        en = content_texts[1] if len(content_texts) > 1 else ''
        jp = content_texts[2] if len(content_texts) > 2 else ''

        # ドリル文スライド（✓ と ✕）
        if content_texts[0] == '✓':
            pos = content_texts[1] if len(content_texts) > 1 else ''
            neg = content_texts[3] if len(content_texts) > 3 else ''
            romaji = f"{pos} / {neg}"
            
            # ドリル文のひらがな化
            kana_text = pos.replace('Watashi wa', 'わたしは')\
                           .replace('Watashi-tachi wa', 'わたしたちは')\
                           .replace('Anata wa', 'あなたは')\
                           .replace('Anata-tachi wa', 'あなたたちは')\
                           .replace('Kare wa', 'かれは')\
                           .replace('Kanojo wa', 'かのじょは')\
                           .replace('Jon', 'ジョン')\
                           .replace('Seeko', 'せいこ')\
                           .replace('Nihon-jin', 'にほんじん')\
                           .replace('Amerika-jin', 'アメリカじん')\
                           .replace('Oosutoraria-jin', 'オーストラリアじん')\
                           .replace('Chuugoku-jin', 'ちゅうごくじん')\
                           .replace('de wa ari-masen', 'ではありません')\
                           .replace('desu', 'です')
            kana = kana_text
            en = 'Substitution Drill'
        elif '?' in romaji:
            # 疑問文・応答文スライド
            q = romaji
            ans_pos = content_texts[2] if len(content_texts) > 2 else ''
            ans_neg = content_texts[4] if len(content_texts) > 4 else ''
            
            kana = q.replace('Anata wa', 'あなたは')\
                    .replace('Jon-san', 'ジョンさん')\
                    .replace('Nihon-jin', 'にほんじん')\
                    .replace('Amerika-jin', 'アメリカじん')\
                    .replace('Oosutoraria-jin', 'オーストラリアじん')\
                    .replace('Chuugoku-jin', 'ちゅうごくじん')\
                    .replace('desu ka?', 'ですか？')
            if ans_pos:
                kana += " " + ans_pos.replace('Hai,', 'はい、').replace('Jon', 'ジョン').replace('Nihon-jin', 'にほんじん').replace('Amerika-jin', 'アメリカじん').replace('Oosutoraria-jin', 'オーストラリアじん').replace('Chuugoku-jin', 'ちゅうごくじん').replace('desu.', 'です。')
            en = 'Question & Answer Drill'
        elif jp:
            kana = jp
        else:
            kana = romaji

        data[i + 1] = {
            'romaji': romaji,
            'kana': kana,
            'en': en
        }
    return data

def main():
    workspace = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    file_001 = os.path.join(workspace, '#001_02_日本語の文字と発音_FC.pptx')
    file_002 = os.path.join(workspace, '#002_02_Lesson1_あいさつ_FC.pptx')
    out_js = os.path.join(workspace, 'js', 'speech-data.js')

    speech_database = {}

    if os.path.exists(file_001):
        print(f"Reading {file_001}...")
        data_001 = extract_001_fc(file_001)
        speech_database["slides/001_02_japanese_writing_fc.pdf"] = data_001
        print(f"Extracted {len(data_001)} items for #001 FC")

    if os.path.exists(file_002):
        print(f"Reading {file_002}...")
        data_002 = extract_002_fc(file_002)
        speech_database["slides/002_02_lesson1_aisatsu_fc.pdf"] = data_002
        print(f"Extracted {len(data_002)} items for #002 FC")

    js_content = f"""/**
 * 日本語音声発音データ辞書 (Web Speech API 用)
 * 自動生成スクリプト: scripts/extract_speech_data.py
 */

const SPEECH_DATA = {json.dumps(speech_database, ensure_ascii=False, indent=2)};
"""

    with open(out_js, 'w', encoding='utf-8') as f:
        f.write(js_content)

    print(f"Successfully generated {out_js} (Total entries: {sum(len(v) for v in speech_database.values())})")

if __name__ == '__main__':
    main()
